"""Import a stalljournal kept in the Google Sheets layout.

Expected sheets (by name):

* ``data``             one row per event; a row with ``ppn`` is a plus row, a row
                       with ``mottagare`` a minus row, rows with neither are herd events
* ``individuppgifter`` animal register and ordered ear tags
* ``ppn-regiser``      other holdings
* ``Flik-B``           holding details as label / value pairs

The spreadsheet mixes events with running totals and has a few rows on the wrong
sheet, so rows are interpreted rather than copied. Everything that was changed or
couldn't be resolved goes into the report.
"""

import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime

from flask import current_app

from . import clock, store
from .constants import EVENT_TYPES
from .db import get_db
from .ids import animal_ref, normalize_ref, parse_tag
from .rules import entry_flags, normalize_ppn, real_vehicle_reg


class ImportRefused(Exception):
    pass


DATA_COLUMNS = {
    "ppn": "ppn",
    "individ-nr": "individual",
    "antal-individer": "count",
    "händelsedatum": "date",
    "identfieringsdatum": "id_date",
    "identifieringsdatum": "id_date",
    "födelseår": "birth_year",
    "leverantör-eller-mor-nr": "supplier",
    "individ-namn": "name",
    "övrigt-plus": "note_plus",
    "mottagare": "recipient",
    "transportör": "driver",
    "fordon": "vehicle",
    "övrigt-minus": "note_minus",
}

TAG_TYPE_WORDS = {"dubbel": "double", "enkel": "single"}

# "födda lamm", but not "vårfödda" (born in spring) in a note about weaning.
BIRTH_WORD = re.compile(r"(?<![a-zåäö])född")
SEX_WORDS = {"tacka": "F", "get": "F", "bagge": "M", "bock": "M"}


# --- cell helpers --------------------------------------------------------------


def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _int(value):
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if text and re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    return None


def _partial_date(value):
    """Full date, or YYYY-MM for values like '2020-04-DD'."""
    full = _date(value)
    if full:
        return full
    text = _text(value)
    if text and re.match(r"\d{4}-\d{2}", text):
        return text[:7]
    return None


def _rows(ws):
    """Rows as dicts keyed by the header row."""
    rows = ws.iter_rows(values_only=True)
    header = [(_text(h) or "").lower() for h in next(rows, [])]
    for number, values in enumerate(rows, start=2):
        yield number, dict(zip(header, values))


# --- the importer --------------------------------------------------------------


class Importer:
    def __init__(self, workbook, filename):
        self.wb = workbook
        self.filename = filename
        self.own_ppn = None
        self.species = "sheep"
        self.entries = []          # (sort key, entry)
        self.herd_events = []
        self.checkpoints = []      # (row, date, count, note)
        self.changed = []          # notes about rows that were interpreted
        self.unresolved = []       # things that need the keeper's attention
        self.animal_tag_types = {}
        self.counts = Counter()

    # holding ----------------------------------------------------------------

    def holding(self):
        if "Flik-B" not in self.wb.sheetnames:
            raise ImportRefused("The workbook has no 'Flik-B' sheet with the holding details.")
        pairs = {}
        for row in self.wb["Flik-B"].iter_rows(values_only=True):
            for i in range(0, len(row) - 1, 3):
                label, value = _text(row[i]), row[i + 1]
                if label and value not in (None, "") and label not in pairs:
                    pairs[label] = value

        ppn = normalize_ppn(_text(pairs.get("Anläggningsnummer")))
        if not ppn:
            raise ImportRefused("Flik-B has no 'Anläggningsnummer'.")
        self.own_ppn = ppn
        art = (_text(pairs.get("Art")) or "").lower()
        self.species = "goat" if art.startswith("get") else "sheep"
        categories = [
            re.split(r"\s+-\s+", _text(pairs[k]))[0]
            for k in ("Kategori-I", "Kategori-II", "Kategori-III") if pairs.get(k)
        ]
        area = pairs.get("Areal")
        registered = _date(pairs.get("Fr.o.m. datum"))
        data = {
            "ppn": ppn,
            "keeper_name": _text(pairs.get("Namn")) or "",
            "address": _text(pairs.get("Adress")),
            "postal_address": _text(pairs.get("Postadress")),
            "county": _text(pairs.get("Län")),
            "municipality": _text(pairs.get("Kommun")),
            "property_designation": _text(pairs.get("Fastighetsbeteckning")),
            "area_ha": float(area) if isinstance(area, (int, float)) else None,
            "housing": _text(pairs.get("Djuren hålls")),
            "species": self.species,
            "max_capacity": _int(pairs.get("Max kapacitet")),
            "registered_from": registered,
            "journal_start": registered,
            "category": "; ".join(categories) or None,
            "phone": _text(pairs.get("Mobilnummer")) or _text(pairs.get("Telefonnummer")),
            "email": _text(pairs.get("e-Postadress")),
        }
        store.save_holding(data, commit=False)
        return data

    # contacts ---------------------------------------------------------------

    def contacts(self):
        if "ppn-regiser" not in self.wb.sheetnames:
            return
        for _, row in _rows(self.wb["ppn-regiser"]):
            ppn = normalize_ppn(_text(row.get("ppn")))
            if not ppn or ppn == self.own_ppn:
                continue
            self._add_contact(ppn, _text(row.get("gård")) or ppn, keeper=_text(row.get("aktör")),
                              phone=_text(row.get("tel-num")))

    def _add_contact(self, ppn, name, kind="farm", keeper=None, phone=None):
        if store.contact_by_ppn(ppn):
            return
        store.save_contact(
            {"ppn": ppn, "name": name, "kind": kind, "keeper": keeper, "phone": phone},
            commit=False,
        )
        self.counts["contacts"] += 1

    # animals and ear tags ---------------------------------------------------

    def animals(self):
        if "individuppgifter" not in self.wb.sheetnames:
            return
        tags = []
        for _, row in _rows(self.wb["individuppgifter"]):
            tag_text = _text(row.get("öronbricka-nr"))
            individual_text = _text(row.get("individ-nr"))
            kind = (_text(row.get("bricka-typ")) or "").lower()
            tag_type = next((t for word, t in TAG_TYPE_WORDS.items() if word in kind), None)
            tag_model = kind.split("-")[0] if kind else None
            colour = _text(row.get("bricka-färg"))
            order_ref = _text(row.get("journal-nr"))
            order_ref = None if order_ref in ("-", None) else order_ref
            species = "goat" if (_text(row.get("djurslag")) or "").lower().startswith("get") else self.species

            # Own tags go into stock whether used or not. 'YYYY' marks the tag year.
            stock = parse_tag((tag_text or "").replace("YYYY", "0000"))
            if stock and stock[0] == self.own_ppn:
                tags.append({
                    "species": species, "individual_number": stock[2],
                    "tag_type": tag_type or "single", "tag_model": tag_model,
                    "tag_colour": colour, "order_ref": order_ref,
                })

            parsed = parse_tag(individual_text)
            if not parsed:
                continue
            ppn, year, number = parsed
            birth_date = _partial_date(row.get("födelsdatum") or row.get("födelsedatum"))
            father = _text(row.get("far-id"))
            data = {
                "species": species,
                "ppn": ppn,
                "individual_number": number,
                "tag_year": _int(row.get("använd-år")) or year,
                "tag_type": tag_type,
                "tag_model": tag_model,
                "tag_colour": colour,
                "name": _text(row.get("individ-namn")),
                "sex": SEX_WORDS.get((_text(row.get("individ-kön")) or "").lower()),
                "birth_date": birth_date,
                "birth_year": int(birth_date[:4]) if birth_date else (_int(row.get("använd-år")) or year),
                "colour": _text(row.get("individ-färg")),
                "breed": _text(row.get("individ-ras")),
                "markings": _text(row.get("kännetecken")),
                "mother_id": normalize_ref(_text(row.get("mor-id"))),
                "mother_name": _text(row.get("mor-namn")),
                "mother_elite_no": _text(row.get("elitlamm-reg-nr-mor")),
                "father_id": normalize_ref(father),
                "father_name": _text(row.get("far-namn")),
                "father_elite_no": _text(row.get("elitlamm-reg-nr-far")),
                "elite_no": _text(row.get("elitlamm-reg-nr")),
                "elite_code": _text(row.get("elitlammskod")),
            }
            store.merge_animal(data, commit=False)
            self.animal_tag_types[(ppn, number)] = tag_type
            self.counts["animals"] += 1
        self.counts["tags"] += store.add_tags(tags, commit=False)

    # journal ----------------------------------------------------------------

    def journal(self):
        if "data" not in self.wb.sheetnames:
            raise ImportRefused("The workbook has no 'data' sheet with the journal rows.")
        arrivals_seen = set()
        borrowed = set()           # animals / groups that arrived as someone else's
        pending_purchases = []     # group purchase rows of 1 waiting for their identity
        unnumbered_tagging = defaultdict(list)

        for number, raw in _rows(self.wb["data"]):
            row = {DATA_COLUMNS[k]: v for k, v in raw.items() if k in DATA_COLUMNS}
            ppn = normalize_ppn(_text(row.get("ppn")))
            individual = _text(row.get("individual"))
            recipient = normalize_ppn(_text(row.get("recipient")))
            note = _text(row.get("note_plus")) or _text(row.get("note_minus")) or ""
            low = note.lower()
            event_date = _date(row.get("date"))
            id_date = _date(row.get("id_date"))
            count = _int(row.get("count"))
            birth_year = _int(row.get("birth_year"))
            supplier_text = _text(row.get("supplier"))
            tag = parse_tag(individual)
            if not any(v not in (None, "") for v in row.values()):
                continue

            common = {
                "species": self.species,
                "driver_name": _text(row.get("driver")),
                "vehicle_reg": _text(row.get("vehicle")),
                "notes": note or None,
                "source": f"import:{self.filename} row {number}",
                "breeding": 0,
                "electronic_tag": 0,
            }

            # Herd events: neither a plus nor a minus row.
            if not ppn and not recipient:
                when = id_date or event_date
                if not when:
                    self.unresolved.append(f"Row {number} ('{note}') has no date and was skipped.")
                    continue
                self.herd_events.append({
                    "event_date": when,
                    "kind": "shearing" if "klipp" in low else "other",
                    "species": self.species,
                    "ppn": tag[0] if tag else None,
                    "individual_number": tag[2] if tag else None,
                    "subject": None if tag else individual,
                    "description": note or None,
                    "source": common["source"],
                })
                continue

            # Minus rows (mottagare set).
            if recipient:
                if tag:
                    animal = {"ppn": tag[0], "individual_number": tag[2], "count": 1}
                else:
                    group_ppn = ppn or normalize_ppn(individual) or self.own_ppn
                    animal = {"ppn": group_ppn, "individual_number": None, "count": count or 1}
                if not event_date:
                    self.unresolved.append(f"Row {number} ('{note}') has no date and was skipped.")
                    continue
                if recipient == self.own_ppn:
                    kind = "return_home"
                    counterparty = normalize_ppn(supplier_text)
                    self.changed.append(
                        f"Row {number}: animals coming back to the holding were on the minus sheet; "
                        "imported as 'Back from temporary stay' (plus)."
                    )
                elif "slakt" in low:
                    kind, counterparty = "slaughter", recipient
                    name = re.search(r"\(([^)]+)\)", note)
                    self._add_contact(recipient, name.group(1) if name else f"Slaughterhouse {recipient}",
                                      kind="slaughterhouse")
                elif "tillfällig" in low:
                    counterparty = recipient
                    key = (animal["ppn"], animal["individual_number"])
                    if key in borrowed and animal["ppn"] == recipient:
                        kind = "loan_return"
                        self.changed.append(
                            f"Row {number}: an animal from {recipient} going back there was imported "
                            "as 'Returned to its owner'."
                        )
                    else:
                        kind = "temp_out"
                elif "död" in low or "dog" in low:
                    kind, counterparty = "death", None
                    common["cause_of_death"] = note
                elif "sål" in low:
                    kind, counterparty = "sale", recipient
                else:
                    kind, counterparty = "other_out", recipient
                self._add_entry(number, kind, event_date, animal, common, counterparty=counterparty)
                continue

            # Plus rows.
            if "ingående balans" in low or "kvar på" in low or (ppn, low) in arrivals_seen:
                self.checkpoints.append((number, event_date, count or 0, note))
                continue

            if tag and not event_date and not id_date:
                match = next((p for p in pending_purchases if p[1]["ppn"] == tag[0]), None)
                if match is None:
                    self.unresolved.append(
                        f"Row {number}: identity {individual} without a date and no matching purchase row."
                    )
                    continue
                pending_purchases.remove(match)
                entry = match[1]
                entry.update({"individual_number": tag[2], "count": 1, "breeding": 1,
                              "birth_year": birth_year or tag[1]})
                self.changed.append(
                    f"Row {number}: identity and birth year of {individual} merged into the purchase "
                    f"on {entry['event_date']} (row {match[0]})."
                )
                continue

            if id_date and not event_date:
                if tag:
                    mother = parse_tag(supplier_text)
                    self._add_entry(number, "tagging", id_date,
                                    {"ppn": tag[0], "individual_number": tag[2], "count": 1}, common,
                                    birth_year=birth_year or tag[1],
                                    mother_id=animal_ref(mother[0], mother[2]) if mother else None,
                                    breeding=1 if self.animal_tag_types.get((tag[0], tag[2])) == "double" else 0)
                else:
                    unnumbered_tagging[id_date].append(number)
                continue

            if not event_date:
                self.unresolved.append(f"Row {number} ('{note}') has no date and was skipped.")
                continue

            if BIRTH_WORD.search(low) or (birth_year and normalize_ppn(supplier_text) == ppn == self.own_ppn):
                males = re.search(r"(\d+)\s*bagg", low)
                females = re.search(r"(\d+)\s*tack", low)
                if tag:
                    self.changed.append(
                        f"Row {number}: {individual} in the individual column is the mother; "
                        "imported as the mother of the litter."
                    )
                self._add_entry(number, "birth", event_date,
                                {"ppn": ppn, "individual_number": None, "count": count or 1}, common,
                                birth_year=birth_year or int(event_date[:4]),
                                mother_id=animal_ref(tag[0], tag[2]) if tag else None,
                                males=int(males.group(1)) if males else None,
                                females=int(females.group(1)) if females else None)
                continue

            supplier = normalize_ppn(supplier_text)
            kind = "purchase" if "köp" in low else "loan_in"
            animal = ({"ppn": tag[0], "individual_number": tag[2], "count": 1} if tag
                      else {"ppn": ppn, "individual_number": None, "count": count or 1})
            entry = self._add_entry(number, kind, event_date, animal, common,
                                    counterparty=supplier, birth_year=birth_year)
            arrivals_seen.add((ppn, low))
            if kind == "loan_in":
                borrowed.add((animal["ppn"], animal["individual_number"]))
            if kind == "purchase" and not tag and (count or 1) == 1:
                pending_purchases.append((number, entry))

        for number, entry in pending_purchases:
            self.unresolved.append(
                f"Row {number}: purchased animal from {entry['ppn']} on {entry['event_date']} "
                "has no individual number."
            )
        for id_date, numbers in unnumbered_tagging.items():
            self.unresolved.append(
                f"Rows {numbers[0]}–{numbers[-1]}: {len(numbers)} lambs ear-tagged on {id_date} without "
                "individual numbers. Not imported: record them with Ear-tag lambs."
            )

    def _add_entry(self, number, kind, event_date, animal, common, **fields):
        entry = {
            **common,
            **animal,
            **fields,
            "event_type": kind,
            "direction": EVENT_TYPES[kind].direction,
            "event_date": event_date,
        }
        self.entries.append(((event_date, number), entry))
        return entry

    # write ------------------------------------------------------------------

    def save(self):
        for _, entry in sorted(self.entries, key=lambda e: e[0]):
            store.insert_entry(entry, created_at=f"{entry['event_date']}T12:00:00", commit=False)
            self.counts[f"entry:{entry['event_type']}"] += 1
        if self.herd_events:
            store.insert_herd_events(self.herd_events, commit=False)

    # report -----------------------------------------------------------------

    def checks(self, today):
        notes = []
        # Running totals from the spreadsheet against the journal balance.
        by_date = defaultdict(list)
        for number, day, count, note in self.checkpoints:
            if day:
                by_date[day].append((number, count, note))
        rows = []
        for day in sorted(by_date):
            stated = sum(c for _, c, _ in by_date[day])
            before = store.balance(before=day).get(self.species, {"present": 0})["present"]
            after = store.balance(as_of=day).get(self.species, {"present": 0})["present"]
            ok = stated in (before, after)
            numbers = ", ".join(str(n) for n, _, _ in by_date[day])
            rows.append(f"| {day} | {numbers} | {stated} | {before} | {after} | {'✓' if ok else '**check**'} |")
        table = []
        if rows:
            table = [
                "Running totals in the spreadsheet compared with the journal (start and end of that day):",
                "",
                "| Date | Rows | Spreadsheet | Journal start of day | Journal end of day | |",
                "|---|---|---|---|---|---|",
                *rows,
            ]

        # Groups of other keepers' animals that never left.
        groups = defaultdict(int)
        for _, entry in self.entries:
            if entry["individual_number"] or not EVENT_TYPES[entry["event_type"]].counted:
                continue
            sign = 1 if entry["direction"] == "in" else -1
            groups[entry["ppn"]] += sign * entry["count"]
        for ppn, left in sorted(groups.items()):
            if ppn != self.own_ppn and left > 0:
                contact = store.contact_by_ppn(ppn)
                label = f"{ppn} ({contact['name']})" if contact else ppn
                notes.append(
                    f"{left} animal(s) with PPN {label} arrived as a group and have no departure. "
                    "If they left, record when and how."
                )

        animals = store.individual_animals()
        unrecorded = sorted(a["ref"] for a in animals.values() if a["status"] == "unrecorded")
        if unrecorded:
            notes.append(
                f"{len(unrecorded)} animal(s) in the register have no journal rows: "
                f"{', '.join(unrecorded)}. Record their ear-tagging (or purchase) in the journal."
            )
        placeholders = sorted({
            e["vehicle_reg"] for _, e in self.entries
            if e["vehicle_reg"] and not real_vehicle_reg(e["vehicle_reg"])
        })
        if placeholders:
            notes.append(
                f"Vehicle registration(s) {', '.join(placeholders)} look like placeholders. "
                "Replace them with the real numbers where you know them."
            )
        missing = sum(
            1 for e in store.list_entries(include_voided=False)
            if "no_transport" in entry_flags(e) and e["vehicle_reg"] not in placeholders
        )
        if missing:
            notes.append(f"{missing} departure(s) have no driver or vehicle registration.")

        dates = sorted(e["event_date"] for _, e in self.entries)
        if dates:
            last = date.fromisoformat(dates[-1])
            if (today - last).days > 180:
                notes.append(
                    f"The journal has no entries after {last.isoformat()}. Record what has happened since "
                    "(departures of the lambs, later births, deaths)."
                )
            missing = [
                y for y in range(int(dates[0][:4]), today.year)
                if not store.get_count(y, self.species)
            ]
            if missing:
                notes.append(
                    f"No December count recorded for {', '.join(map(str, missing))}. If you reported "
                    "counts to Jordbruksverket, add them under December count."
                )
        return table, notes


def import_workbook(path, dry_run=False, replace=False):
    from openpyxl import load_workbook
    import warnings

    db = get_db()
    if store.entry_count() and not replace:
        raise ImportRefused(
            "The journal already has entries. Use --replace to delete the current journal, animals, "
            "ear tags, contacts, herd log and December counts and import from scratch."
        )
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, data_only=True, read_only=False)

    filename = os.path.basename(path)
    importer = Importer(workbook, filename)
    try:
        if replace:
            for table in ("entry", "animal", "ear_tag", "contact", "herd_event", "annual_count"):
                db.execute(f"DELETE FROM {table}")
        holding = importer.holding()
        importer.contacts()
        importer.animals()
        importer.journal()
        importer.save()
        today = clock.today()
        table, notes = importer.checks(today)
        stock = store.balance().get(importer.species, {"present": 0, "away": 0})
        report = _report(importer, holding, table, notes, stock, dry_run, replace)
        if dry_run:
            db.rollback()
        else:
            db.commit()
    except Exception:
        db.rollback()
        raise

    report_path = os.path.join(
        os.path.dirname(os.path.abspath(current_app.config["DATABASE"])),
        f"import-report-{clock.now().strftime('%Y%m%d-%H%M%S')}{'-dry-run' if dry_run else ''}.md",
    )
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report)
    return report + f"\nReport saved to {report_path}\n"


def _report(importer, holding, table, notes, stock, dry_run, replace):
    c = importer.counts
    entry_counts = {k.split(":", 1)[1]: v for k, v in c.items() if k.startswith("entry:")}
    lines = [
        f"# Import of {importer.filename}",
        "",
        "**Dry run: nothing was saved.**" if dry_run else "Saved to the database.",
        "",
    ]
    if replace:
        lines += ["The previous journal, animals, ear tags, contacts, herd log and December counts were deleted first.", ""]
    lines += [
        "## Imported",
        "",
        f"- Holding: PPN {holding['ppn']}, {holding['keeper_name']}"
        f"{', ' + holding['category'] if holding['category'] else ''}",
        f"- Contacts: {c['contacts']}",
        f"- Animals in the register: {c['animals']}",
        f"- Ear tags in stock: {c['tags']}",
        f"- Journal rows: {sum(entry_counts.values())}"
        + (" (" + ", ".join(f"{EVENT_TYPES[k].label.lower()} {v}" for k, v in sorted(entry_counts.items())) + ")"
           if entry_counts else ""),
        f"- Herd log entries: {len(importer.herd_events)}",
        f"- Running-total rows used only as checks: {len(importer.checkpoints)}",
        "",
    ]
    if importer.changed:
        lines += ["## Interpreted on the way", "", *[f"- {n}" for n in importer.changed], ""]
    attention = importer.unresolved + notes
    if attention or table:
        lines += ["## Needs your attention", ""]
        lines += [f"- {n}" for n in attention]
        if table:
            lines += ["", *table]
        lines.append("")
    lines += [
        "## Balance after import",
        "",
        f"- On the holding according to the journal: {stock['present']}",
        f"- Temporarily away: {stock['away']}",
        "",
    ]
    return "\n".join(lines)
