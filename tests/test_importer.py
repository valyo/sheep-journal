"""The spreadsheet importer, against a small made-up workbook in the same layout."""

from datetime import datetime as dt

import pytest
from openpyxl import Workbook

from sheep_journal import store
from sheep_journal.importer import ImportRefused, import_workbook

DATA_HEADER = [
    "gård", "ppn", "individ-nr", "antal-individer", "händelsedatum", "identfieringsdatum",
    "födelseår", "leverantör-eller-mor-nr", "individ-namn", "övrigt-plus", "mottagare",
    "transportör", "fordon", "övrigt-minus",
]

ANIMAL_HEADER = [
    "individ-nr", "gård", "journal-nr", "öronbricka-nr", "bricka-typ", "bricka-färg", "djurslag",
    "använd-år", "individ-namn", "födelsdatum", "individ-kön", "individ-färg", "individ-ras",
    "mor-id", "far-id",
]


def data_row(**cells):
    return [cells.get(h) for h in DATA_HEADER]


@pytest.fixture
def workbook_path(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)

    b = wb.create_sheet("Flik-B")
    b.append(["anläggningsinformation", "data", None, "aktörer", "data", None, "arter", "data"])
    b.append(["Anläggningsnummer", "SE 11111", None, "Namn", "Test Keeper", None, "Art", "Får"])
    b.append(["Kommun", "Testby", None, "Fr.o.m. datum", dt(2023, 2, 28), None, "Max kapacitet", 10])
    b.append(["Areal", 5, None, "Mobilnummer", "070000", None, "Kategori-I", "Avelsbesättning - 2 - fr.o.m. 2023-02-28"])

    r = wb.create_sheet("ppn-regiser")
    r.append(["ppn-num", "gård", "ppn", "aktör", "tel-num"])
    r.append([1, "Hemma", "SE11111", "Test Keeper", "070000"])
    r.append([2, "Grannby", "SE22222", "Granne", "070111"])

    a = wb.create_sheet("individuppgifter")
    a.append(ANIMAL_HEADER)
    a.append(["SE33333-2019-19011", "Ute", "-", "SE33333-2019-19011", "micro", "grön", "får", 2019,
              "Stina", dt(2019, 3, 16), "tacka", "svart", "klövsjöfår", None, None])
    a.append(["SE11111-2023-00051", "Hemma", "2023/1", "SE11111-YYYY-00051", "micro-dubbel", "blå",
              "får", 2023, "Kanel", dt(2023, 5, 7), "tacka", "vit", "blandras",
              "SE33333-2019-19011", "Grannens bagge"])
    a.append(["SE11111-2023-00001", "Hemma", "2023/1", "SE11111-YYYY-00001", "micro-enkel", "blå",
              "får", 2023, "Rulle", dt(2023, 5, 7), "bagge", "svart", "blandras",
              "SE33333-2019-19011", None])
    a.append([None, "Hemma", "2023/1", "SE11111-YYYY-00052", "micro-dubbel", "blå", "får"])
    a.append(["SE11111-2024-00005", "Hemma", "2023/1", "SE11111-YYYY-00005", "micro-enkel", "blå",
              "får", 2024])

    d = wb.create_sheet("data")
    d.append(DATA_HEADER)
    rows = [
        data_row(ppn="SE33333", **{"antal-individer": 1, "händelsedatum": dt(2023, 4, 13),
                 "leverantör-eller-mor-nr": "SE22222", "övrigt-plus": "köpt tacka från Grannby",
                 "transportör": "T. Keeper", "fordon": "ABC 123"}),
        data_row(ppn="SE33333", **{"individ-nr": "SE33333-2019-19011", "födelseår": 2019,
                 "övrigt-plus": "reg. i journalen köpta tackor identitet och födelseår"}),
        data_row(ppn="SE11111", **{"antal-individer": 1, "händelsedatum": dt(2023, 5, 7),
                 "övrigt-plus": "får på gården (ingående balans)"}),
        data_row(ppn="SE11111", **{"individ-nr": "SE33333-2019-19011", "antal-individer": 2,
                 "händelsedatum": dt(2023, 5, 7), "födelseår": 2023,
                 "leverantör-eller-mor-nr": "SE11111", "övrigt-plus": "födda lamm (1 bagge och 1 tacka)"}),
        data_row(ppn="SE11111", **{"individ-nr": "SE11111-2023-00051", "identfieringsdatum": dt(2023, 9, 8),
                 "födelseår": 2023, "leverantör-eller-mor-nr": "SE33333-2019-19011",
                 "övrigt-plus": "öronmarkering"}),
        data_row(ppn="SE11111", **{"individ-nr": "SE11111-2023-00001", "identfieringsdatum": dt(2023, 9, 8),
                 "födelseår": 2023, "leverantör-eller-mor-nr": "SE33333-2019-19011",
                 "övrigt-plus": "öronmarkering"}),
        data_row(ppn="SE11111", **{"identfieringsdatum": dt(2024, 5, 1), "födelseår": 2024,
                 "övrigt-plus": "öronmarkering"}),
        data_row(**{"individ-nr": "SE33333-2019-19011", "händelsedatum": dt(2023, 9, 10),
                 "leverantör-eller-mor-nr": "SE11111", "mottagare": "SE22222",
                 "transportör": "T. Keeper", "fordon": "ABC 123", "övrigt-minus": "tillfälligt förflyttas tackan"}),
        data_row(ppn="SE22222", **{"antal-individer": 3, "händelsedatum": dt(2023, 9, 10), "födelseår": 2023,
                 "leverantör-eller-mor-nr": "SE22222", "övrigt-plus": "inkommit på produktionsplatsen",
                 "transportör": "T. Keeper", "fordon": "ABC 123"}),
        data_row(ppn="SE22222", **{"antal-individer": 3, "händelsedatum": dt(2023, 10, 10),
                 "övrigt-plus": "inkommit på produktionsplatsen"}),
        data_row(**{"individ-nr": "SE11111-2023-00001", "händelsedatum": dt(2023, 11, 28),
                 "mottagare": "SE099", "transportör": "T. Keeper", "fordon": "XXX 000",
                 "övrigt-minus": "slakt (TEST SLAKTERI AB)"}),
        data_row(**{"individ-nr": "SE33333-2019-19011", "händelsedatum": dt(2024, 3, 3),
                 "leverantör-eller-mor-nr": "SE22222", "mottagare": "SE11111",
                 "transportör": "T. Keeper", "fordon": "XXX 000", "övrigt-minus": "tackan hämtas tillbaka"}),
        data_row(**{"individ-nr": "SE33333-2019-19011", "identfieringsdatum": dt(2024, 5, 22),
                 "individ-namn": "Stina", "övrigt-plus": "fårklipp"}),
    ]
    for row in rows:
        d.append(row)

    path = tmp_path / "journal.xlsx"
    wb.save(path)
    return str(path)


def test_import(app, workbook_path):
    with app.app_context():
        report = import_workbook(workbook_path)
        holding = store.get_holding()
        assert holding["ppn"] == "11111" and holding["category"] == "Avelsbesättning"
        assert holding["max_capacity"] == 10 and holding["registered_from"] == "2023-02-28"

        assert store.contact_by_ppn("22222")["name"] == "Grannby"
        assert store.contact_by_ppn("099")["kind"] == "slaughterhouse"
        assert store.contact_by_ppn("11111") is None

        kinds = [(e["event_type"], e["ppn"], e["individual_number"], e["count"])
                 for e in store.list_entries()]
        assert kinds == [
            ("purchase", "33333", "19011", 1),
            ("birth", "11111", None, 2),
            ("tagging", "11111", "00051", 1),
            ("tagging", "11111", "00001", 1),
            ("temp_out", "33333", "19011", 1),
            ("loan_in", "22222", None, 3),
            ("slaughter", "11111", "00001", 1),
            ("return_home", "33333", "19011", 1),
        ]
        entries = {e["event_type"] + (e["individual_number"] or ""): e for e in store.list_entries()}
        assert entries["birth"]["mother_id"] == "33333-19011"
        assert (entries["birth"]["males"], entries["birth"]["females"]) == (1, 1)
        assert entries["tagging00051"]["breeding"] == 1
        assert entries["tagging00001"]["breeding"] == 0
        assert entries["purchase19011"]["breeding"] == 1
        assert entries["return_home19011"]["counterparty"] == "22222"
        assert entries["purchase19011"]["source"].startswith("import:journal.xlsx row 2")

        assert store.balance()["sheep"] == {"present": 5, "away": 0}
        animals = store.individual_animals()
        assert animals[("sheep", "11111", "00051")]["name"] == "Kanel"
        assert animals[("sheep", "11111", "00051")]["father_id"] == "Grannens bagge"
        assert animals[("sheep", "11111", "00001")]["status"] == "gone"
        assert animals[("sheep", "11111", "00005")]["status"] == "unrecorded"
        assert [t["individual_number"] for t in store.available_tags("sheep")] == ["00052"]
        assert store.list_herd_events()[0]["kind"] == "shearing"

    assert "running-total rows used only as checks: 2" in report.lower()
    assert "minus sheet" in report
    assert "without individual numbers" in report
    assert "3 animal(s) with PPN 22222 (Grannby)" in report
    assert "11111-00005" in report
    assert "XXX 000" in report


def test_dry_run_saves_nothing(app, workbook_path):
    with app.app_context():
        report = import_workbook(workbook_path, dry_run=True)
        assert "Dry run" in report
        assert store.entry_count() == 0 and store.get_holding() is None


def test_refuses_when_journal_has_entries(app, workbook_path):
    with app.app_context():
        import_workbook(workbook_path)
        with pytest.raises(ImportRefused):
            import_workbook(workbook_path)
        import_workbook(workbook_path, replace=True)
        assert store.entry_count() == 8
